#!/usr/bin/env python3
"""Build manuscript/Mystery_of_the_Trinity.docx — final interior layout.

Renders the manuscript markdown into a 4.25 x 6.875 in (trim size) portrait
Word document with the final front matter (half-title, Note on the Text, full
title page, blank page, and a live Table of Contents field), chapter-level
page breaks, a centered 8 pt page-number footer (suppressed on the first
page), and the Garamond style set: body 12 pt, Chapter headings 16 pt bold,
Section headings 12.5 pt bold italic, prayers italic indented. See the
project notes for the style specification.
"""

import re
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

BASE_DIR = Path(__file__).resolve().parent.parent
MD_PATH = BASE_DIR / "manuscript" / "Mystery_of_the_Trinity.md"
DOCX_PATH = BASE_DIR / "manuscript" / "Mystery_of_the_Trinity.docx"

BYLINE_MARKER = "by Fr. Peter Gruber, C.O."
EDITION_LINE = "Working manuscript draft \u00b7 2026"
TOC_NOTE = "(Select and update this table in Word to fill in page numbers.)"

# Final interior trim: 4.25 x 6.875 in, portrait, matching the book PDF's
# margins and footer position.
PAGE_W = Inches(4.25)
PAGE_H = Inches(6.875)
MARGIN_LR = Inches(0.5)
MARGIN_TOP = Inches(0.8)
MARGIN_BOTTOM = Inches(0.65)
FOOTER_DIST = Inches(0.4)

# Live TOC field: collect the custom Chapter style at level 1 and the Section
# style at level 2. \h hyperlinks the entries, \z omits page numbers in web
# view, \u uses the paragraphs' outline levels where present.
TOC_INSTR = ' TOC \\t "Chapter,1,Section,2" \\h \\z \\u '

TOC_LINK_RE = re.compile(r"^(\s*)- \[(.+?)\]\(#[^)]*\)$")
NUM_ITEM_RE = re.compile(r"^\d+\.\s")


# ---------------------------------------------------------------------------
# Styles
# ---------------------------------------------------------------------------
def set_garamond(style):
    """Set Garamond on a style: font name plus explicit w:ascii / w:hAnsi."""
    style.font.name = "Garamond"
    rpr = style.element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    rf.set(qn("w:ascii"), "Garamond")
    rf.set(qn("w:hAnsi"), "Garamond")


def make_styles(doc):
    normal = doc.styles["Normal"]
    set_garamond(normal)
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    def add_para_style(name):
        try:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = normal
        except ValueError:
            style = doc.styles[name]
            if style.type != WD_STYLE_TYPE.PARAGRAPH:
                # The default template ships "Book Title" as a *character*
                # style; drop it and create a paragraph style with that name.
                style.element.getparent().remove(style.element)
                style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
                style.base_style = normal
        set_garamond(style)
        style.font.size = Pt(12)  # explicit; subclasses override below
        return style

    half_title = add_para_style("Half Title")
    half_title.font.size = Pt(22)
    half_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    book_title = add_para_style("Book Title")
    book_title.font.size = Pt(26)
    book_title.font.bold = True
    book_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    book_title.paragraph_format.keep_with_next = True

    chapter = add_para_style("Chapter")
    chapter.font.size = Pt(16)
    chapter.font.bold = True
    chapter.paragraph_format.space_before = Pt(12)
    chapter.paragraph_format.space_after = Pt(6)
    chapter.paragraph_format.keep_with_next = True

    section = add_para_style("Section")
    section.font.size = Pt(12.5)
    section.font.bold = True
    section.font.italic = True
    section.paragraph_format.space_before = Pt(10)
    section.paragraph_format.space_after = Pt(4)
    section.paragraph_format.keep_with_next = True

    prayer = add_para_style("Prayer")
    prayer.font.italic = True
    prayer.paragraph_format.left_indent = Inches(0.5)
    prayer.paragraph_format.space_after = Pt(6)

    attribution = add_para_style("Attribution")
    attribution.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    attribution.paragraph_format.space_after = Pt(6)

    return {
        "Half Title": half_title,
        "Book Title": book_title,
        "Chapter": chapter,
        "Section": section,
        "Prayer": prayer,
        "Attribution": attribution,
    }


# ---------------------------------------------------------------------------
# Inline markup
# ---------------------------------------------------------------------------
def add_runs(paragraph, text, base_bold=False, base_italic=False):
    """Add runs for *text*, splitting on ** (bold) first, then * (italic)."""
    bold_on = False
    for segment in text.split("**"):
        italic_on = False
        for part in segment.split("*"):
            if part:
                run = paragraph.add_run(part)
                run.bold = base_bold or bold_on
                run.italic = base_italic or italic_on
                run.font.name = "Garamond"
            italic_on = not italic_on
        bold_on = not bold_on


# ---------------------------------------------------------------------------
# OOXML fields (page number, live table of contents)
# ---------------------------------------------------------------------------
def _append_field(run, instr_text, cached_text, dirty=False):
    """Attach a complex field (begin/instr/separate/result/end) to *run*."""
    r = run._r
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    if dirty:
        begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instr_text
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    cached = OxmlElement("w:t")
    cached.text = cached_text
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, cached, end):
        r.append(element)


def add_page_number_field(paragraph):
    run = paragraph.add_run()
    run.font.name = "Garamond"
    run.font.size = Pt(8)
    _append_field(run, " PAGE ", "1")


def add_toc_field(paragraph):
    run = paragraph.add_run()
    run.font.name = "Garamond"
    run.font.size = Pt(12)
    _append_field(
        run,
        TOC_INSTR,
        "Right-click here and choose 'Update Field' to build the table of contents.",
        dirty=True,
    )


# ---------------------------------------------------------------------------
# Section geometry and footer
# ---------------------------------------------------------------------------
def make_section(doc):
    section = doc.sections[0]
    section.page_width = PAGE_W
    section.page_height = PAGE_H
    section.left_margin = MARGIN_LR
    section.right_margin = MARGIN_LR
    section.top_margin = MARGIN_TOP
    section.bottom_margin = MARGIN_BOTTOM
    section.footer_distance = FOOTER_DIST
    # First page (half-title) carries no page number.
    section.different_first_page_header_footer = True

    footer_para = section.footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_para.paragraph_format.space_after = Pt(0)
    add_page_number_field(footer_para)
    # section.first_page_footer is left empty on purpose.


# ---------------------------------------------------------------------------
# Markdown subset parser
# ---------------------------------------------------------------------------
def parse_markdown(text):
    lines = text.split("\n")
    blocks = []
    current = []

    def flush_paragraph():
        if current:
            blocks.append(("para", " ".join(current)))
            current.clear()

    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()
        if stripped == "" or stripped == "---":
            flush_paragraph()
        elif stripped.startswith("# "):
            flush_paragraph()
            blocks.append(("h1", stripped[2:].strip()))
        elif stripped.startswith("## "):
            flush_paragraph()
            blocks.append(("h2", stripped[3:].strip()))
        elif stripped.startswith("### "):
            flush_paragraph()
            blocks.append(("h3", stripped[4:].strip()))
        elif line.startswith(">"):
            flush_paragraph()
            raw = []
            while i < n and lines[i].startswith(">"):
                rest = lines[i][1:]
                raw.append(rest[1:] if rest.startswith(" ") else rest)
                i += 1
            blocks.append(("blockquote", raw))
            continue
        else:
            toc_match = TOC_LINK_RE.match(line)
            if toc_match:
                flush_paragraph()
                level = 2 if len(toc_match.group(1)) >= 2 else 1
                blocks.append(("toc", toc_match.group(2).strip(), level))
            elif NUM_ITEM_RE.match(stripped):
                flush_paragraph()
                blocks.append(("listitem", stripped))
            else:
                current.append(stripped)
        i += 1
    flush_paragraph()
    return blocks


def is_wrapped_italic(text):
    """True when the line is a standalone paragraph entirely wrapped in *...*."""
    return (
        len(text) > 1
        and text.startswith("*")
        and text.endswith("*")
        and text.count("*") == 2
    )


def split_blockquote(lines):
    """Split blockquote raw lines into paragraphs on blank lines."""
    groups, current = [], []
    for line in lines:
        if line == "":
            if current:
                groups.append(" ".join(current))
                current = []
        else:
            current.append(line)
    if current:
        groups.append(" ".join(current))
    return groups


# ---------------------------------------------------------------------------
# Rendering
# ---------------------------------------------------------------------------
def render_front_matter(doc, styles, title, byline, note_paras):
    # Page 1 — half-title: book title only, centered.
    paragraph = doc.add_paragraph(style=styles["Half Title"])
    add_runs(paragraph, title)

    # Page 2 — Note on the Text.
    paragraph = doc.add_paragraph(style=styles["Chapter"])
    paragraph.paragraph_format.page_break_before = True
    add_runs(paragraph, "Note on the Text")
    for text in note_paras:
        paragraph = doc.add_paragraph(style="Normal")
        add_runs(paragraph, text)

    # Page 3 — full title page: title, author, edition line.
    paragraph = doc.add_paragraph(style=styles["Book Title"])
    paragraph.paragraph_format.page_break_before = True
    add_runs(paragraph, title)
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(12)
    add_runs(paragraph, byline)
    paragraph = doc.add_paragraph(style="Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(paragraph, EDITION_LINE)

    # Page 4 — blank (an empty paragraph whose page break opens the page).
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.page_break_before = True

    # Page 5 — Table of Contents heading + live TOC field + note.
    paragraph = doc.add_paragraph(style=styles["Chapter"])
    paragraph.paragraph_format.page_break_before = True
    add_runs(paragraph, "Table of Contents")
    paragraph = doc.add_paragraph(style="Normal")
    add_toc_field(paragraph)
    paragraph = doc.add_paragraph(style="Normal")
    add_runs(paragraph, TOC_NOTE)


def render_body(doc, styles, blocks):
    for block in blocks:
        kind = block[0]
        if kind == "h1":
            paragraph = doc.add_paragraph(style=styles["Chapter"])
            paragraph.paragraph_format.page_break_before = True
            add_runs(paragraph, block[1])
        elif kind in ("h2", "h3"):
            paragraph = doc.add_paragraph(style=styles["Section"])
            add_runs(paragraph, block[1])
        elif kind == "para":
            text = block[1]
            if is_wrapped_italic(text):
                inner = text[1:-1].strip()
                if inner.startswith("—"):
                    # italic attribution, e.g. "*— Prayer of St. Elizabeth ...*"
                    paragraph = doc.add_paragraph(style=styles["Attribution"])
                    add_runs(paragraph, inner, base_italic=True)
                else:
                    # standalone prayer paragraph
                    paragraph = doc.add_paragraph(style=styles["Prayer"])
                    add_runs(paragraph, inner, base_italic=True)
            elif text.startswith("—"):
                paragraph = doc.add_paragraph(style=styles["Attribution"])
                add_runs(paragraph, text)
            else:
                paragraph = doc.add_paragraph(style="Normal")
                add_runs(paragraph, text)
        elif kind == "blockquote":
            for quote in split_blockquote(block[1]):
                if quote.startswith("—"):
                    paragraph = doc.add_paragraph(style=styles["Attribution"])
                    add_runs(paragraph, quote, base_italic=True)
                else:
                    paragraph = doc.add_paragraph(style=styles["Prayer"])
                    add_runs(paragraph, quote, base_italic=True)
        elif kind == "toc":
            # Static TOC links are replaced by the live field; ignore any
            # that leak into the body slice.
            continue
        else:  # listitem
            paragraph = doc.add_paragraph(style="Normal")
            add_runs(paragraph, block[1])


def find_heading_index(blocks, heading_text):
    """Index of the first h1 block with *heading_text*, or raise."""
    for i, block in enumerate(blocks):
        if block[0] == "h1" and block[1] == heading_text:
            return i
    raise ValueError(f"heading not found: {heading_text!r}")


def split_front_matter(blocks):
    """Split the title, byline, Note-on-the-Text text, and body blocks."""
    assert blocks[0][0] == "h1", "manuscript must start with the book title"
    title = blocks[0][1]
    assert blocks[1][0] == "para" and blocks[1][1] == BYLINE_MARKER, blocks[1]
    byline = blocks[1][1]

    note_start = find_heading_index(blocks, "Note on the Text")
    toc_start = find_heading_index(blocks, "Table of Contents")
    note_paras = [b[1] for b in blocks[note_start + 1 : toc_start] if b[0] == "para"]

    last_toc = max(i for i, b in enumerate(blocks) if b[0] == "toc")
    return title, byline, note_paras, blocks[last_toc + 1 :]


def report(doc, title):
    style_counts = Counter(p.style.name for p in doc.paragraphs)
    summary = ", ".join(
        f"{name}={count}" for name, count in sorted(style_counts.items())
    )
    page_breaks = sum(
        1 for p in doc.paragraphs if p.paragraph_format.page_break_before
    )
    section = doc.sections[0]
    print(f"Wrote {DOCX_PATH.relative_to(BASE_DIR)}")
    print(f"  paragraphs: {len(doc.paragraphs)}  pageBreakBefore: {page_breaks}")
    print(f"  title: {title!r}")
    print(
        "  page: "
        f"{section.page_width} x {section.page_height} EMU "
        f"margins L/R {section.left_margin}/{section.right_margin} "
        f"T {section.top_margin} B {section.bottom_margin} "
        f"footer_dist {section.footer_distance}"
    )
    print(f"  styles: {summary}")
    print(f"  file size: {DOCX_PATH.stat().st_size} bytes")


def main():
    text = MD_PATH.read_text(encoding="utf-8")
    title, byline, note_paras, body_blocks = split_front_matter(parse_markdown(text))

    doc = Document()
    styles = make_styles(doc)
    make_section(doc)
    render_front_matter(doc, styles, title, byline, note_paras)
    render_body(doc, styles, body_blocks)

    doc.core_properties.title = title
    doc.save(DOCX_PATH)
    report(doc, title)


if __name__ == "__main__":
    main()
